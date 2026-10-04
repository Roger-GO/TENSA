"""Integration tests: which buses a loaded case gives no rated voltage.

ANDES fills in 110 kV for a bus whose ``Vn`` the case leaves out or zeroes, and
the topology names those buses in ``buses_without_vn`` so a client never reads
the fill-in as a voltage base. Each case format keeps a different trace of the
original value, so every one is loaded for real here, from ANDES's bundled cases
and from copies of them with the ``Vn`` of some buses taken out.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from tensa.core.wrapper import Wrapper

pytestmark = pytest.mark.integration


def _cases() -> Path:
    pytest.importorskip("andes")
    import andes

    return Path(andes.__file__).parent / "cases"


def _load(path: Path) -> Wrapper:
    wrapper = Wrapper()
    wrapper.load_case(path)
    return wrapper


def _listed(wrapper: Wrapper) -> list[int | str]:
    return wrapper.topology_snapshot().buses_without_vn


@pytest.mark.parametrize(
    "case",
    ["ieee14/ieee14.raw", "ieee14/ieee14.json", "ieee14/ieee14_full.xlsx"],
)
def test_a_case_that_rates_every_bus_lists_none(case: str) -> None:
    assert _listed(_load(_cases() / case)) == []


# ---- MATPOWER ---------------------------------------------------------------


def test_matpower_case_14_gives_no_base_kv_so_every_bus_is_listed() -> None:
    # Its baseKV column is 0 on every row, and the reader swaps in 110 for each.
    wrapper = _load(_cases() / "matpower" / "case14.m")
    assert _listed(wrapper) == list(range(1, 15))


def test_matpower_lists_only_the_buses_with_a_zero_base_kv(tmp_path: Path) -> None:
    text = (_cases() / "matpower" / "case14.m").read_text(encoding="utf-8")
    # baseKV is the tenth column of the bus matrix: give buses 1 to 3 one.
    for bus in (1, 2, 3):
        text = re.sub(
            rf"^(\t{bus}(?:\t[^\t\n]+){{8}})\t0\t", r"\1\t135\t", text, count=1, flags=re.M
        )
    case = tmp_path / "case14_partly_rated.m"
    case.write_text(text, encoding="utf-8")

    wrapper = _load(case)
    assert _listed(wrapper) == list(range(4, 15))
    assert sorted(set(wrapper._ss.Bus.Vn.v)) == [110, 135]  # noqa: SLF001


# ---- PSS/E RAW --------------------------------------------------------------


def test_raw_lists_the_buses_with_a_zero_basekv(tmp_path: Path) -> None:
    lines = (_cases() / "ieee14" / "ieee14.raw").read_text(encoding="latin-1").split("\n")
    for number in (1, 6):
        at = next(i for i, line in enumerate(lines) if line.lstrip().startswith(f"{number},'BUS"))
        fields = lines[at].split(",")
        fields[2] = "  0.0000"
        lines[at] = ",".join(fields)
    case = tmp_path / "ieee14_unrated.raw"
    case.write_text("\n".join(lines), encoding="latin-1")

    assert _listed(_load(case)) == [1, 6]


# ---- json -------------------------------------------------------------------


def test_json_lists_the_buses_with_a_missing_null_or_zero_vn(tmp_path: Path) -> None:
    data = json.loads((_cases() / "ieee14" / "ieee14.json").read_text(encoding="utf-8"))
    del data["Bus"][0]["Vn"]
    data["Bus"][1]["Vn"] = None
    data["Bus"][2]["Vn"] = 0
    case = tmp_path / "ieee14_unrated.json"
    case.write_text(json.dumps(data), encoding="utf-8")

    assert _listed(_load(case)) == [1, 2, 3]


def test_json_without_a_vn_anywhere_lists_every_bus(tmp_path: Path) -> None:
    data = json.loads((_cases() / "ieee14" / "ieee14.json").read_text(encoding="utf-8"))
    for row in data["Bus"]:
        row.pop("Vn")
    case = tmp_path / "ieee14_no_vn.json"
    case.write_text(json.dumps(data), encoding="utf-8")

    assert _listed(_load(case)) == list(range(1, 15))


# ---- xlsx -------------------------------------------------------------------


def _xlsx_vn_column(sheet: Worksheet) -> int:
    return next(cell.column for cell in sheet[1] if cell.value == "Vn")


def test_xlsx_lists_the_buses_with_a_blank_or_zero_vn(tmp_path: Path) -> None:
    book = load_workbook(_cases() / "ieee14" / "ieee14_full.xlsx")
    sheet = book["Bus"]
    column = _xlsx_vn_column(sheet)
    sheet.cell(row=2, column=column).value = None
    sheet.cell(row=3, column=column).value = 0
    case = tmp_path / "ieee14_unrated.xlsx"
    book.save(case)

    assert _listed(_load(case)) == [1, 2]


def test_xlsx_without_a_vn_column_lists_every_bus(tmp_path: Path) -> None:
    book = load_workbook(_cases() / "ieee14" / "ieee14_full.xlsx")
    sheet = book["Bus"]
    sheet.delete_cols(_xlsx_vn_column(sheet))
    case = tmp_path / "ieee14_no_vn.xlsx"
    book.save(case)

    assert _listed(_load(case)) == list(range(1, 15))


# ---- what happens to the list afterwards ------------------------------------


def test_the_listing_survives_setup_and_a_power_flow() -> None:
    # ANDES clears its own record of the corrections when it sets up.
    wrapper = _load(_cases() / "matpower" / "case14.m")
    wrapper.run_pflow()
    assert wrapper.topology_snapshot().state == "committed"
    assert _listed(wrapper) == list(range(1, 15))


def test_a_bus_given_a_vn_after_the_load_is_no_longer_listed() -> None:
    wrapper = _load(_cases() / "matpower" / "case14.m")
    wrapper.edit_element("Bus", 1, {"Vn": 138.0})
    assert _listed(wrapper) == list(range(2, 15))


def test_a_reload_lists_the_same_buses() -> None:
    wrapper = _load(_cases() / "matpower" / "case14.m")
    wrapper.edit_element("Bus", 1, {"Vn": 138.0})
    wrapper.reload_case()
    assert _listed(wrapper) == list(range(1, 15))
