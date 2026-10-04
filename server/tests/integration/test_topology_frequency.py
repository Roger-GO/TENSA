"""Integration tests: where the topology's ``freq_hz`` comes from, per format.

The base frequency is ``System.config.freq``. A PSS/E RAW file sets it in its
header, and an xlsx or json file sets it in its ``_config`` section, which ANDES
merges into the configuration. A MATPOWER file has no frequency to give, so its
System keeps ANDES's default of 60 Hz.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tensa.core.wrapper import Wrapper

pytestmark = pytest.mark.integration


def _cases() -> Path:
    pytest.importorskip("andes")
    import andes

    return Path(andes.__file__).parent / "cases"


def _freq_hz(path: Path) -> float | None:
    return Wrapper().load_case(path).freq_hz


def test_a_matpower_case_reads_andes_default_frequency() -> None:
    assert _freq_hz(_cases() / "matpower" / "case14.m") == 60.0


def test_a_json_case_reads_the_frequency_of_its_config_section(tmp_path: Path) -> None:
    data = json.loads((_cases() / "ieee14" / "ieee14.json").read_text(encoding="utf-8"))
    assert "_config" not in data
    assert _freq_hz(_cases() / "ieee14" / "ieee14.json") == 60.0

    data["_config"] = [{"section": "System", "key": "freq", "value": 50.0}]
    case = tmp_path / "ieee14_50hz.json"
    case.write_text(json.dumps(data), encoding="utf-8")
    assert _freq_hz(case) == 50.0


def test_an_xlsx_case_reads_the_frequency_of_its_config_sheet(tmp_path: Path) -> None:
    book = load_workbook(_cases() / "ieee14" / "ieee14_full.xlsx")
    assert "_config" not in book.sheetnames
    sheet = book.create_sheet("_config")
    sheet.append(["section", "key", "value"])
    sheet.append(["System", "freq", 50])
    case = tmp_path / "ieee14_50hz.xlsx"
    book.save(case)

    assert _freq_hz(case) == 50.0
