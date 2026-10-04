"""Integration tests: the topology's ``events``, per source.

A case file can define timed events of its own. ANDES's bundled ``kundur_full.xlsx``
has a ``Toggle`` that trips ``Line_8`` at 2 s, so a time-domain run of it swings with
nothing scheduled by the client, and a client that says "nothing disturbs this run"
is wrong. The topology names those events, and the ones a snapshot restore replays.
What the client commits itself is not among them.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tensa.core.case_events import CaseEvent
from tensa.core.disturbance import FaultSpec
from tensa.core.wrapper import Wrapper

pytestmark = pytest.mark.integration

_KUNDUR_TRIP = CaseEvent(
    source="case", kind="toggle", t=2.0, name="Toggler_1", model="Line", dev_idx="Line_8"
)


def _cases() -> Path:
    pytest.importorskip("andes")
    import andes

    return Path(andes.__file__).parent / "cases"


def test_kundur_full_lists_the_line_trip_its_file_defines() -> None:
    assert Wrapper().load_case(_cases() / "kundur" / "kundur_full.xlsx").events == [_KUNDUR_TRIP]


def test_a_case_without_events_lists_none() -> None:
    assert Wrapper().load_case(_cases() / "ieee14" / "ieee14.raw").events == []
    assert Wrapper().load_case(_cases() / "matpower" / "case14.m").events == []


def test_an_xlsx_lists_each_kind_and_leaves_out_what_cannot_act(tmp_path: Path) -> None:
    book = load_workbook(_cases() / "ieee14" / "ieee14_full.xlsx")
    fault = book.create_sheet("Fault")
    fault.append(["idx", "u", "name", "bus", "tf", "tc", "xf", "rf"])
    fault.append(["F1", 1, "Fault_a", 9, 1.0, 1.1, 0.05, 0])
    fault.append(["F2", 0, "Fault_off", 9, 1.5, 1.6, 0.05, 0])  # switched off
    fault.append(["F3", 1, "Fault_open", 4, 3.0, -1, 0.05, 0])  # never cleared
    alter = book.create_sheet("Alter")
    alter.append(["idx", "u", "name", "t", "model", "dev", "src", "method", "amount"])
    alter.append(["A1", 1, "Alter_1", 2.5, "PQ", "PQ_1", "Ppf", "*", 1.2])
    toggle = book.create_sheet("Toggler")
    toggle.append(["idx", "u", "name", "model", "dev", "t"])
    toggle.append(["T1", 1, "Trip_1", "Line", "Line_3", 4.0])
    case = tmp_path / "events.xlsx"
    book.save(case)

    events = Wrapper().load_case(case).events

    assert [(e.kind, e.name, e.t, e.tc) for e in events] == [
        ("fault", "Fault_a", 1.0, 1.1),
        ("alter", "Alter_1", 2.5, None),
        ("fault", "Fault_open", 3.0, None),
        ("toggle", "Trip_1", 4.0, None),
    ]
    assert events[0].dev_idx == 9 and events[0].model == "Bus"
    assert (events[1].model, events[1].dev_idx, events[1].src) == ("PQ", "PQ_1", "Ppf")
    assert (events[1].method, events[1].amount) == ("*", 1.2)
    assert all(e.source == "case" for e in events)


def test_the_listing_survives_setup_a_power_flow_and_a_reload() -> None:
    w = Wrapper()
    w.load_case(_cases() / "kundur" / "kundur_full.xlsx")
    w.run_pflow()

    committed = w.topology_snapshot()
    assert committed.state == "committed"
    assert committed.events == [_KUNDUR_TRIP]
    assert w.reload_case().events == [_KUNDUR_TRIP]


def test_what_the_client_commits_is_not_listed() -> None:
    w = Wrapper()
    w.load_case(_cases() / "kundur" / "kundur_full.xlsx")

    w.add_disturbance(FaultSpec(bus_idx=3, tf=1.0, tc=1.1))

    assert w.topology_snapshot().events == [_KUNDUR_TRIP]


def test_loading_another_case_drops_the_events_of_the_first() -> None:
    w = Wrapper()
    w.load_case(_cases() / "kundur" / "kundur_full.xlsx")
    assert w.load_case(_cases() / "ieee14" / "ieee14.raw").events == []


def test_a_restored_snapshot_lists_the_disturbances_it_replayed(tmp_path: Path) -> None:
    workspace = tmp_path / "ws"
    workspace.mkdir(mode=0o700)
    case = workspace / "ieee14.raw"
    shutil.copy2(_cases() / "ieee14" / "ieee14.raw", case)
    w = Wrapper(workspace=workspace, session_id="case-events")
    w.load_case(case)
    w.add_disturbance(FaultSpec(bus_idx=5, tf=1.0, tc=1.1))
    w.save_snapshot("snap")

    w.load_case(case)
    assert w.topology_snapshot().events == []

    w.restore_snapshot("snap")
    assert w.topology_snapshot().events == [
        CaseEvent(source="restored", kind="fault", t=1.0, tc=1.1, model="Bus", dev_idx=5)
    ]
    # A new System has none of them.
    assert w.reload_case().events == []
