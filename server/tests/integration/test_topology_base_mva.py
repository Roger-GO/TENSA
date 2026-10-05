"""Integration tests: where the topology's ``base_mva`` comes from, per format.

The system MVA base is ``System.config.mva``. A PSS/E RAW file sets it in its
header, a MATPOWER file in ``baseMVA``, and an xlsx or json file in its
``_config`` section, which ANDES merges into the configuration. A case that sets
none, and a blank system, keeps ANDES's default of 100.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tensa.core.wrapper import Wrapper

pytestmark = pytest.mark.integration


def _cases() -> Path:
    pytest.importorskip("andes")
    import andes

    return Path(andes.__file__).parent / "cases"


def _base_mva(path: Path) -> float | None:
    return Wrapper().load_case(path).base_mva


def test_a_blank_system_has_andes_default_base() -> None:
    assert Wrapper().create_blank().base_mva == 100.0


def test_a_raw_case_reads_the_base_of_its_header(tmp_path: Path) -> None:
    source = _cases() / "ieee14" / "ieee14.raw"
    assert _base_mva(source) == 100.0

    # The base is the second field of the RAW header line.
    header, rest = source.read_text(encoding="latin-1").split("\n", 1)
    assert "100.00" in header
    case = tmp_path / "ieee14_200mva.raw"
    case.write_text(header.replace("100.00", "200.00", 1) + "\n" + rest, encoding="latin-1")
    assert _base_mva(case) == 200.0


def test_a_matpower_case_reads_its_base_mva(tmp_path: Path) -> None:
    source = _cases() / "matpower" / "case14.m"
    assert _base_mva(source) == 100.0

    text = source.read_text(encoding="utf-8")
    assert "mpc.baseMVA = 100;" in text
    case = tmp_path / "case14_200mva.m"
    case.write_text(text.replace("mpc.baseMVA = 100;", "mpc.baseMVA = 200;"), encoding="utf-8")
    assert _base_mva(case) == 200.0


def test_a_json_case_reads_the_base_of_its_config_section(tmp_path: Path) -> None:
    data = json.loads((_cases() / "ieee14" / "ieee14.json").read_text(encoding="utf-8"))
    assert "_config" not in data
    assert _base_mva(_cases() / "ieee14" / "ieee14.json") == 100.0

    data["_config"] = [{"section": "System", "key": "mva", "value": 250.0}]
    case = tmp_path / "ieee14_250mva.json"
    case.write_text(json.dumps(data), encoding="utf-8")
    assert _base_mva(case) == 250.0


def test_an_xlsx_case_reads_the_base_of_its_config_sheet(tmp_path: Path) -> None:
    book = load_workbook(_cases() / "ieee14" / "ieee14_full.xlsx")
    assert "_config" not in book.sheetnames
    sheet = book.create_sheet("_config")
    sheet.append(["section", "key", "value"])
    sheet.append(["System", "mva", 250])
    case = tmp_path / "ieee14_250mva.xlsx"
    book.save(case)

    assert _base_mva(case) == 250.0

